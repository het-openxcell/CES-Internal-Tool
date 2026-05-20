import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

_HEADER_BG = "FF111827"
_HEADER_FILL = PatternFill("solid", fgColor=_HEADER_BG)
_HEADER_FONT = Font(color="FFFFFFFF", bold=True)


class MasterExcelExportService:
    def generate(self, occurrences_with_source: list[dict]) -> io.BytesIO:
        wb = Workbook()
        ws = wb.active
        ws.title = "Occurrences"

        headers = ["Well Name", "Surface Location", "Type", "Section", "mMD", "Density", "Notes", "DDR Source"]
        for col, h in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.fill = _HEADER_FILL
            cell.font = _HEADER_FONT

        if not occurrences_with_source:
            ws.cell(row=2, column=1, value="No occurrences found")
        else:
            for row_idx, row in enumerate(occurrences_with_source, start=2):
                ws.cell(row=row_idx, column=1, value=row.get("well_name"))
                ws.cell(row=row_idx, column=2, value=row.get("surface_location"))
                ws.cell(row=row_idx, column=3, value=row.get("type"))
                ws.cell(row=row_idx, column=4, value=row.get("section"))
                ws.cell(row=row_idx, column=5, value=row.get("mmd"))
                ws.cell(row=row_idx, column=6, value=row.get("density"))
                ws.cell(row=row_idx, column=7, value=row.get("notes"))
                ws.cell(row=row_idx, column=8, value=row.get("ddr_source"))

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf
