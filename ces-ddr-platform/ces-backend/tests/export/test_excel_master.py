from unittest.mock import MagicMock

import openpyxl
import pytest

from src.services.export.excel_master import MasterExcelExportService


def _make_row(i: int) -> dict:
    return {
        "well_name": f"Well {i}",
        "surface_location": f"Surface {i}",
        "type": "Drilling",
        "section": "Main",
        "mmd": float(1000 + i),
        "density": 1.5,
        "notes": f"Note {i}",
        "ddr_source": f"DDR_{i}.pdf",
    }


def test_generate_empty_list_returns_header_plus_no_data_row():
    service = MasterExcelExportService()
    buf = service.generate([])
    wb = openpyxl.load_workbook(buf)
    ws = wb["Occurrences"]
    assert ws.cell(row=2, column=1).value == "No occurrences found"


def test_generate_two_rows_returns_header_plus_two_data_rows():
    service = MasterExcelExportService()
    rows = [_make_row(0), _make_row(1)]
    buf = service.generate(rows)
    wb = openpyxl.load_workbook(buf)
    ws = wb["Occurrences"]
    assert ws.max_row == 3


def test_sheet_has_eight_columns_including_ddr_source():
    service = MasterExcelExportService()
    buf = service.generate([_make_row(0)])
    wb = openpyxl.load_workbook(buf)
    ws = wb["Occurrences"]
    headers = [ws.cell(row=1, column=c).value for c in range(1, 9)]
    assert "DDR Source" in headers
    assert len(headers) == 8


@pytest.mark.anyio
async def test_export_master_route_returns_200(monkeypatch):
    from fastapi.testclient import TestClient

    from src.main import backend_app
    from src.securities.authorizations.jwt_authentication import jwt_authentication

    async def fake_get_all(self, **kwargs):
        return [_make_row(0)]

    monkeypatch.setattr(
        "src.repository.crud.occurrence.OccurrenceCRUDRepository.get_all_with_ddr_source",
        fake_get_all,
    )
    backend_app.dependency_overrides[jwt_authentication] = lambda: MagicMock(id="user-1")
    client = TestClient(backend_app)
    response = client.get("/api/export/master")
    backend_app.dependency_overrides.clear()
    assert response.status_code == 200
    assert "spreadsheetml" in response.headers["content-type"]


@pytest.mark.anyio
async def test_export_master_filtered_filename(monkeypatch):
    from fastapi.testclient import TestClient

    from src.main import backend_app
    from src.securities.authorizations.jwt_authentication import jwt_authentication

    async def fake_get_all(self, **kwargs):
        return []

    monkeypatch.setattr(
        "src.repository.crud.occurrence.OccurrenceCRUDRepository.get_all_with_ddr_source",
        fake_get_all,
    )
    backend_app.dependency_overrides[jwt_authentication] = lambda: MagicMock(id="user-1")
    client = TestClient(backend_app)
    response = client.get("/api/export/master?ddr_id=some-id")
    backend_app.dependency_overrides.clear()
    assert "occurrences_filtered.xlsx" in response.headers["content-disposition"]
