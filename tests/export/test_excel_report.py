import io
from dataclasses import dataclass
from unittest.mock import MagicMock

import openpyxl
import pytest

from src.services.export.excel_report import DDRExcelExportService


@dataclass
class FakeOccurrence:
    well_name: str | None
    surface_location: str | None
    type: str
    section: str | None
    mmd: float | None
    density: float | None
    notes: str | None


@dataclass
class FakeCorrection:
    field_name: str
    original_value: str
    corrected_value: str
    reason: str
    user_id: str
    created_at: int


def _make_occurrence(i: int) -> FakeOccurrence:
    return FakeOccurrence(
        well_name=f"Well {i}",
        surface_location=f"Surface {i}",
        type="Drilling",
        section="Main",
        mmd=float(1000 + i),
        density=1.5,
        notes=f"Note {i}",
    )


def _make_correction(i: int) -> FakeCorrection:
    return FakeCorrection(
        field_name="type",
        original_value="Ream",
        corrected_value="Back Ream",
        reason="Misclassified",
        user_id=f"user-{i}",
        created_at=1716000000 + i,
    )


def test_generate_returns_valid_xlsx():
    service = DDRExcelExportService()
    occs = [_make_occurrence(i) for i in range(3)]
    buf = service.generate(occs, [], "DDR_TEST.pdf")
    assert isinstance(buf, io.BytesIO)
    wb = openpyxl.load_workbook(buf)
    assert "Occurrences" in wb.sheetnames


def test_sheet1_header_style():
    service = DDRExcelExportService()
    buf = service.generate([_make_occurrence(0)], [], "DDR_TEST.pdf")
    wb = openpyxl.load_workbook(buf)
    ws = wb["Occurrences"]
    header_cell = ws.cell(row=1, column=1)
    assert header_cell.fill.fgColor.rgb == "FF111827"
    assert header_cell.font.color.rgb == "FFFFFFFF"


def test_sheet2_present_when_corrections():
    service = DDRExcelExportService()
    corrections = [_make_correction(1)]
    buf = service.generate([_make_occurrence(0)], corrections, "DDR_TEST.pdf")
    wb = openpyxl.load_workbook(buf)
    assert "Edit History" in wb.sheetnames


def test_sheet2_absent_when_no_corrections():
    service = DDRExcelExportService()
    buf = service.generate([_make_occurrence(0)], [], "DDR_TEST.pdf")
    wb = openpyxl.load_workbook(buf)
    assert "Edit History" not in wb.sheetnames


def test_occurrences_sheet_has_correct_data():
    service = DDRExcelExportService()
    occ = _make_occurrence(1)
    buf = service.generate([occ], [], "DDR_TEST.pdf")
    wb = openpyxl.load_workbook(buf)
    ws = wb["Occurrences"]
    assert ws.cell(row=2, column=1).value == "Well 1"
    assert ws.cell(row=2, column=3).value == "Drilling"


@pytest.mark.anyio
async def test_export_ddr_route_404_for_unknown_ddr(monkeypatch):
    from fastapi.testclient import TestClient

    from src.main import backend_app

    async def fake_read_by_id(self, id):
        return None

    monkeypatch.setattr("src.repository.crud.ddr.DDRCRUDRepository.read_by_id", fake_read_by_id)

    from src.securities.authorizations.jwt_authentication import jwt_authentication
    backend_app.dependency_overrides[jwt_authentication] = lambda: MagicMock(id="user-1")

    client = TestClient(backend_app)
    response = client.get("/api/export/ddr/nonexistent-id")
    backend_app.dependency_overrides.clear()
    assert response.status_code == 404
    assert response.json()["code"] == "NOT_FOUND"
