import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

_HEADER_BG = "FF111827"
_HEADER_FILL = PatternFill("solid", fgColor=_HEADER_BG)
_HEADER_FONT = Font(color="FFFFFFFF", bold=True)


class DDRExcelExportService:
    def generate(self, occurrences: list, corrections: list, ddr_source_label: str) -> io.BytesIO:
        wb = Workbook()
        ws = wb.active
        ws.title = "Occurrences"

        occ_headers = ["Well Name", "Surface Location", "Type", "Section", "mMD", "Density", "Notes"]
        for col, h in enumerate(occ_headers, start=1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.fill = _HEADER_FILL
            cell.font = _HEADER_FONT

        for row_idx, occ in enumerate(occurrences, start=2):
            ws.cell(row=row_idx, column=1, value=occ.well_name)
            ws.cell(row=row_idx, column=2, value=occ.surface_location)
            ws.cell(row=row_idx, column=3, value=occ.type)
            ws.cell(row=row_idx, column=4, value=occ.section)
            ws.cell(row=row_idx, column=5, value=occ.mmd)
            ws.cell(row=row_idx, column=6, value=occ.density)
            ws.cell(row=row_idx, column=7, value=occ.notes)

        if corrections:
            ws2 = wb.create_sheet(title="Edit History")
            edit_headers = ["Field", "Original Value", "Corrected Value", "Reason", "User", "Timestamp", "DDR Source"]
            for col, h in enumerate(edit_headers, start=1):
                cell = ws2.cell(row=1, column=col, value=h)
                cell.fill = _HEADER_FILL
                cell.font = _HEADER_FONT
            for row_idx, c in enumerate(corrections, start=2):
                ws2.cell(row=row_idx, column=1, value=c.field_name)
                ws2.cell(row=row_idx, column=2, value=c.original_value)
                ws2.cell(row=row_idx, column=3, value=c.corrected_value)
                ws2.cell(row=row_idx, column=4, value=c.reason)
                ws2.cell(row=row_idx, column=5, value=c.user_id)
                ws2.cell(row=row_idx, column=6, value=c.created_at)
                ws2.cell(row=row_idx, column=7, value=ddr_source_label)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf
